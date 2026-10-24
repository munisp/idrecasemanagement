"""The QA text of a generated letter keeps the letter's layout: paragraphs, address lines, tables in place."""
import io
import os
import unittest

os.environ.setdefault("MINIO_ENDPOINT", "localhost:9000")
from docx import Document  # noqa: E402

import main  # noqa: E402


def letter() -> bytes:
    d = Document()
    for ln in ("October 18, 2026", "Agency for Health Care Administration", "2727 Mahan Drive", "RE: External Review of Denied Claims"):
        d.add_paragraph(ln)
    d.add_paragraph("Dear Ms. Rooks:")
    d.add_paragraph("Capitol Bridge is an organization that contracts with the Florida Agency for Health Care Administration "
                    "to independently review claim disputes.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Claim", "Amount"
    t.cell(1, 0).text, t.cell(1, 1).text = "C1", "$3,000.00"
    d.add_paragraph("Sincerely,")
    d.add_paragraph("Capitol Bridge")
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


class Preview(unittest.TestCase):
    def test_layout_kept(self):
        s = main.docx_preview(letter())
        self.assertIn("October 18, 2026\nAgency for Health Care Administration\n2727 Mahan Drive\nRE: External Review of Denied Claims", s)
        self.assertIn("RE: External Review of Denied Claims\n\nDear Ms. Rooks:\n\nCapitol Bridge is an organization", s)
        self.assertIn("claim disputes.\n\nClaim | Amount\nC1 | $3,000.00\n\nSincerely,\nCapitol Bridge", s)
        self.assertNotIn("  ", s)

    def test_not_cut_at_4000(self):
        d = Document()
        for i in range(80):
            d.add_paragraph(f"Paragraph {i} " + "word " * 20)
        b = io.BytesIO()
        d.save(b)
        self.assertIn("Paragraph 79", main.docx_preview(b.getvalue()))


if __name__ == "__main__":
    unittest.main()
