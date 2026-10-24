"""The office's letters: every template fills without a blank left, sections and table rows follow the case, and
the fields the letters say are worked out from the case as the office writes them.
   python -m unittest test_office_letters"""
import glob
import io
import os
import unittest
from datetime import date

from docx import Document

import builtin_letters as bl
import docxfill
import office_letters as ol

HERE = os.path.dirname(os.path.abspath(__file__))


def text_of(docx_bytes: bytes) -> str:
    d = Document(io.BytesIO(docx_bytes))
    return "\n".join(docxfill.ptext(p) for p in docxfill.paragraphs(d))


def ctx(outcome="OVERTURNED", claims=2, paid=None, report=True):
    rep = {"report_type": "CLAIM_DISPUTE", "outcome": outcome, "awarded_cents": 125000, "already_paid_cents": paid,
           "percent_awarded": None, "reviewers": ["attorney", "coding consultant"], "disallowed": [],
           "sections": {"claim_summary": "This case concerns two (2) claims.", "dispute_issue": "The plan denied the claims.",
                        "provider_position": "The provider says the codes are right.", "analysis": "The codes are right.",
                        "coding_review": "that the codes billed were supported."},
           "claims": [{"claim_number": "C1", "cpt": "99285", "dos": "08/14/2026", "billed_cents": 90000, "paid_cents": 10000,
                       "ucr_cents": 60000, "reimbursement_cents": 50000}]}
    return {"tenant": "fl", "case_id": "x", "case_number": "FL26-021354", "details": {},
            "dates": {"date_case_accepted": "2026-10-02", "plan_notified_at": "2026-10-03"},
            "disputed_cents": 250000, "num_claims": claims, "num_patients": 1,
            "claims": [("C1", "99285", 90000, 10000, date(2026, 8, 14), "Ann", "Lee"),
                       ("C2", "99284", 50000, 0, date(2026, 8, 15), "Ann", "Lee")][:claims],
            "opened_at": None,
            "provider": {"name": "Gulf Coast Clinic", "contact": "Dr Lee", "title": "", "email": "lee@gulf.test",
                         "address": ["1 Main St", "Tampa, FL 33601"]},
            "plan": {"name": "Sunshine Health", "contact": "", "title": "", "email": "plan@sun.test", "address": []},
            "letters": {"org_name": "Capitol Bridge LLC", "org_short": "Capitol Bridge", "invoice": {}, "remit": {}, "agency": {}},
            "report": rep if report else None,
            "invoices": [{"party": "HEALTH_PLAN", "kind": "FULL_REVIEW", "invoice_no": "FL26-021354", "amount_cents": 60000,
                          "status": "OPEN", "due_date": None}],
            "eligibility": None, "rfi": None, "today": date(2026, 10, 23)}


class Templates(unittest.TestCase):
    def test_every_template_fills_without_a_blank(self):
        for path in sorted(glob.glob(os.path.join(HERE, "assets", "office", "*.docx"))):
            raw = open(path, "rb").read()
            fields = {f: "VALUE" for f in docxfill.fields_in(raw)}
            fields.update({"coding": True, "coding_rows": [{"dos": "D", "code": "C", "billed": "B", "allowed": "A",
                                                            "medicare200": "M", "additional": "X"}], "physician": True})
            out = text_of(docxfill.fill(raw, fields))
            for bad in ("{", "[insert", "Month Day, Year", "MM/DD/YYYY", "XX"):
                self.assertNotIn(bad, out, f"{os.path.basename(path)} still has {bad!r}")

    def test_sections_and_rows_follow_the_case(self):
        raw = open(os.path.join(HERE, "assets", "office", "rationale_partial.docx"), "rb").read()
        base = {f: "V" for f in docxfill.fields_in(raw)}
        full = text_of(docxfill.fill(raw, {**base, "coding": True, "physician": True,
                                           "coding_rows": [{"dos": "08/14/2026"}, {"dos": "08/15/2026"}]}))
        self.assertIn("Independent Physician Review", full)
        self.assertIn("Coding Consultant Review", full)
        self.assertEqual(full.count("08/1"), 2)  # one table row per priced claim
        bare = text_of(docxfill.fill(raw, {**base, "coding": False, "physician": False, "coding_rows": []}))
        self.assertNotIn("Independent Physician Review", bare)
        self.assertNotIn("Coding Consultant Review", bare)
        self.assertNotIn("Date of Service", bare)  # the table went with its section
        self.assertIn("Based on the foregoing", bare)

    def test_an_empty_address_line_leaves_no_line(self):
        raw = open(os.path.join(HERE, "assets", "office", "acceptance.docx"), "rb").read()
        f = {k: "V" for k in docxfill.fields_in(raw)}
        f["provider_address1"] = ""
        d = Document(io.BytesIO(docxfill.fill(raw, f)))
        lines = [docxfill.ptext(p._p) for p in d.paragraphs]
        self.assertFalse(any(x.replace(" ", "").strip() == "" and " " in x for x in lines))


class Fields(unittest.TestCase):
    def test_common(self):
        f = ol._common(ctx())
        self.assertEqual(f["today"], "October 23, 2026")
        self.assertEqual(f["dates_of_service"], "Multiple (2)")
        self.assertEqual(f["member_names"], "Ann Lee")
        self.assertEqual(f["provider_attn"], "ATTN: Dr Lee")
        self.assertEqual(f["plan_salutation"], "Sunshine Health")  # no contact: the organization
        self.assertEqual((f["provider_address1"], f["provider_city_state_zip"]), ("1 Main St", "Tampa, FL 33601"))
        self.assertEqual(f["plan_due_date"], "October 18, 2026")
        one = ol._common(ctx(claims=1))
        self.assertEqual(one["dates_of_service"], "08/14/2026")
        self.assertEqual(one["claims_that_comprise"], "The claim that comprises this case was")

    def test_decision(self):
        c = ctx()
        f = ol._decision_fields(c, ol._common(c))
        self.assertTrue(f["determination_summary"].startswith("the Health Plan’s denial") or "overturned" in f["determination_summary"])
        self.assertFalse(f["determination_summary"].endswith("."))
        self.assertEqual(f["case_determination"], "should be overturned and the amount of $1,250.00 be awarded to the Provider.")
        self.assertEqual(f["coding_review_paragraph"], "The Capitol Bridge coding consultant indicated that the codes billed were supported.")
        self.assertEqual(f["coding_rows"][0]["medicare200"], "$600.00")
        self.assertIn(("b", "Provider’s Position:"), f["rationale"])
        self.assertEqual(f["reviewers_phrase"], "an attorney and a coding consultant")
        p = ctx(outcome="PARTIALLY_OVERTURNED", paid=25000)
        pf = ol._decision_fields(p, ol._common(p))
        self.assertEqual((pf["award_total"], pf["already_paid"], pf["award_additional"]), ("$1,250.00", "$250.00", "$1,000.00"))

    def test_template_choice(self):
        ex = {"case_outcome": "", "internal_status": ""}
        self.assertEqual(ol.template_for("final_order_rationale", ctx(), ex), "rationale")
        pasted = ctx()
        pasted["report"]["sections"]["full_report"] = "The whole report, pasted."
        self.assertIsNone(ol.template_for("final_order_rationale", pasted, ex))  # printed as written instead
        self.assertEqual(ol.template_for("final_order_rationale", ctx(outcome="PARTIALLY_OVERTURNED"), ex), "rationale_partial")
        self.assertEqual(ol.template_for("final_order_rationale", ctx(report=False), {"case_outcome": "Default Award", "internal_status": ""}),
                         "rationale_default")
        self.assertEqual(ol.template_for("invoice_health_plan", ctx(), ex), "invoice_full")
        split = ctx()
        split["invoices"].append({"party": "PROVIDER", "kind": "FULL_REVIEW", "invoice_no": "FL26-021354", "amount_cents": 40000,
                                  "status": "OPEN", "due_date": None})
        self.assertEqual(ol.template_for("invoice_health_plan", split, ex), "invoice_split")

    def test_invoice(self):
        c = ctx()
        f, extra = ol._invoice_fields(c, {"late_fee": {"cents": 5000, "after_days": 35, "every_days": 30}}, "HEALTH_PLAN", ol._common(c))
        self.assertEqual((f["review_cost"], f["total_late1"], f["total_late2"]), ("$600.00", "$650.00", "$700.00"))
        self.assertEqual((f["late_fee"], f["late_after"], f["late_after2"]), ("$50", "35", "65"))
        self.assertEqual(f["remit_bank"], "First Business Bank")  # the office's own until Finance sets one
        self.assertEqual(f["responsible_party"], "Plan")
        self.assertTrue(extra["invoice_number"].startswith("FL26-021354"))

    def test_partial_acceptance_needs_marked_claims(self):
        c = ctx()
        with self.assertRaises(bl.LetterDataMissing):
            ol._stage_fields("partial_acceptance_letter", c, {"claim_status": [("C1", "SUBMITTED")]}, ol._common(c))
        f = ol._stage_fields("partial_acceptance_letter", c, {"claim_status": [("C1", "SUBMITTED"), ("C2", "INELIGIBLE"),
                                                                                ("C3", "SUBMITTED")]}, ol._common(c))
        self.assertEqual((f["accepted_claims"], f["ineligible_claims"]), ("C1 and C3", "C2"))


if __name__ == "__main__":
    unittest.main()
