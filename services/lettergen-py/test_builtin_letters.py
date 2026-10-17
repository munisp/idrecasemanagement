"""Built-in letters: the decision wording, the invoice arithmetic, and the refusals when a case lacks what a letter
must say. Builds from a hand-made case context -- no database, no MinIO.
   python -m unittest test_builtin_letters"""
import io
import unittest
from datetime import date

from docx import Document

import builtin_letters as bl

LETTERS = {
    "org_name": "Capitol Bridge LLC", "org_short": "Capitol Bridge",
    "program_title": "FLORIDA PROVIDER AND HEALTH PLAN CLAIM DISPUTE RESOLUTION PROGRAM",
    "contractor_line": "Florida AHCA Claims Dispute Resolution Contractor", "state": "Florida",
    "agency_name": "Agency for Health Care Administration", "statute": "Chapter 408.7057, Florida Statutes",
    "contact_email": "office@example.test", "contact_phone": "1-800-000-0000",
    "standard_of_review": "Capitol Bridge has interpreted the rights and responsibilities of the parties.",
    "agency": {"salutation": "Dear Ms. Example:", "lines": ["Agency for Health Care Administration", "Tallahassee, FL"]},
    "invoice": {"health_plan_suffix": "H", "late_fee_cents": 5000, "late_after_days": 35, "late_every_days": 30,
                "rule_full": "FL Admin Code R 59A-12.030(4)(a)", "rule_split": "FL Admin Code R 59A-12.030(4)(b)",
                "payment_note": "Capitol Bridge only accepts ACH payment for FL AHCA dispute invoices."},
    "remit": {"payee": "Capitol Bridge LLC", "bank": "Example Bank", "routing": "021000021", "account": "000123456789",
              "remittance_email": "remit@example.test"},
}


def ctx(**over):
    c = {
        "tenant": "fl", "case_id": "c1", "case_number": "FL26-000216", "details": {"line_of_business": "Medicaid"},
        "dates": {"date_case_accepted": "2026-02-02"}, "disputed_cents": 5_000_000, "num_claims": 7, "num_patients": 5,
        "claims": [], "provider": {"name": "Gulf Coast Clinic", "contact": "Pat Doe", "email": "pat@clinic.test",
                                   "address": ["1 Main St", "Miami, FL 33101"]},
        "plan": {"name": "Bay Health Plan", "contact": "", "email": "claims@bay.test", "address": []},
        "letters": LETTERS, "report": None, "invoices": [], "eligibility": None, "rfi": None, "today": date(2026, 9, 1),
    }
    c.update(over)
    return c


def report(**over):
    r = {"report_type": "CLAIM_DISPUTE", "outcome": "OVERTURNED", "awarded_cents": 4_000_000,
         "already_paid_cents": 1_000_000, "percent_awarded": None, "sections": {}, "claims": [], "disallowed": [],
         "reviewers": []}
    r.update(over)
    return r


def words(letter: bl.Letter) -> str:
    doc = Document(io.BytesIO(letter.save()))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


class DecisionWording(unittest.TestCase):
    def test_overturned_with_balance(self):
        s = bl.recommendation(ctx(), report())
        self.assertIn("be overturned and the amount of $40,000.00 be awarded to the Provider", s)
        self.assertIn("of which $10,000.00 was already paid", s)
        self.assertIn("the balance of $30,000.00 is owed to the Provider", s)

    def test_partially_overturned(self):
        s = bl.recommendation(ctx(), report(outcome="PARTIALLY_OVERTURNED"))
        self.assertIn("recommends that the Health Plan’s denial of payment for these services be partially overturned", s)
        s = bl.recommendation(ctx(), report(outcome="PARTIALLY_OVERTURNED"), verb="has determined that")
        self.assertNotIn("that that", s)

    def test_upheld_awards_the_plan(self):
        s = bl.recommendation(ctx(), report(outcome="UPHELD", awarded_cents=0, already_paid_cents=None))
        self.assertIn("be upheld and the amount of $0.00 be awarded to the Health Plan", s)

    def test_report_sections_and_tables(self):
        r = report(outcome="PARTIALLY_OVERTURNED", sections={
            "provider_position": "The Provider says so.", "plan_position": "The plan says otherwise.",
            "physician_review": "Medically necessary.", "coding_review": "Bundled charges remain denied."},
            claims=[{"claim_number": "C1", "cpt": "99285", "dos": "2025-05-01", "zip": "33101", "billed_cents": 300000,
                     "ucr_cents": 250000, "lesser_cents": 250000, "paid_cents": 100000, "reimbursement_cents": 150000}],
            disallowed=[{"reason": "Routine nursing", "description": "Vent day 1", "revenue_code": "0410",
                         "quantity": "-1", "amount_cents": 10000},
                        {"reason": "Routine nursing", "description": "Vent wean", "revenue_code": "0410",
                         "quantity": "-19", "amount_cents": 20000}])
        w = words(bl.build_decision_report(ctx(report=r, disputed_cents=4_347_826), None))
        for want in ("CAPITOL BRIDGE LLC.", "CASE NO: FL26-000216", "CLAIM NO: C1", "Provider’s Position:",
                     "Health Plan’s Position:", "Independent Physician Review:", "Coding Consultant Review:",
                     "UCR Pricing", "$1,500.00", "-20", "$300.00 of disallowed charges will remain denied",
                     "or 92% of the disputed charge", "On February 2, 2026, Capitol Bridge notified Gulf Coast Clinic"):
            self.assertIn(want, w)

    def test_iro_variant(self):
        r = report(report_type="IRO", outcome="OVERTURNED", awarded_cents=None, sections={
            "iro_number": "00051", "date_received": "2026-03-11", "decision_date": "2026-03-15",
            "determination_reason": "the requested service is not experimental or investigational",
            "evidence": "Ref one\nRef two", "professional_review": "Not experimental.",
            "binding_statute": "the state's external review statute"})
        w = words(bl.build_decision_report(ctx(report=r), None))
        for want in ("Decision Summary:", "DATE RECEIVED: March 11, 2026", "DECISION DATE: March 15, 2026",
                     "Evidence Supporting the Determination:", "Ref two",
                     "be overturned because the requested service is not experimental or investigational.",
                     "This decision is final and binding on the insurance carrier and the covered person"):
            self.assertIn(want, w)

    def test_cover_letter(self):
        w = words(bl.build_cover_letter(ctx(report=report(reviewers=["attorney", "physician", "coding consultant"])), None))
        for want in ("RE: External Review of Denied Claims", "The amount in controversy is $50,000.00",
                     "Capitol Bridge LLC. has determined that", "Dear Ms. Example:",
                     "reviewed by an attorney, physician and coding consultant reviewers",
                     "Enclosed with this letter is the analysis and basis for decision"):
            self.assertIn(want, w)

    def test_no_outcome_is_refused(self):
        with self.assertRaises(bl.LetterDataMissing):
            bl.build_decision_report(ctx(report=report(outcome=None)), None)
        with self.assertRaises(bl.LetterDataMissing):
            bl.build_cover_letter(ctx(report=None), None)


class Invoices(unittest.TestCase):
    def inv(self, party, cents):
        return {"party": party, "kind": "FULL_REVIEW", "invoice_no": "FL26-000216", "amount_cents": cents,
                "status": "OPEN", "due_date": None}

    def test_split_invoice(self):
        c = ctx(invoices=[self.inv("HEALTH_PLAN", 322092), self.inv("PROVIDER", 28008)])
        d, extra = bl.build_invoice(c, "HEALTH_PLAN", None)
        w = words(d)
        self.assertEqual(extra["invoice_number"], "FL26-000216H")
        for want in ("Health Plan is responsible for 92% of the review cost", "R 59A-12.030(4)(b)", "$3,220.92",
                     "TOTAL if paid after 35 days: $3,270.92", "TOTAL if paid after 65 days: $3,320.92",
                     "09/01/2026", "See Final", "Bay Health Plan", "021000021", "remit@example.test"):
            self.assertIn(want, w)

    def test_full_invoice_to_provider(self):
        d, extra = bl.build_invoice(ctx(invoices=[self.inv("PROVIDER", 106200)]), "PROVIDER", None)
        w = words(d)
        self.assertEqual(extra["invoice_number"], "FL26-000216")
        for want in ("Provider is responsible for 100% of the review cost", "R 59A-12.030(4)(a)",
                     "TOTAL if paid after 35 days: $1,112.00", "TOTAL if paid after 65 days: $1,162.00", "ATTN: Pat Doe"):
            self.assertIn(want, w)

    def test_missing_invoice_or_bank_is_refused(self):
        with self.assertRaises(bl.LetterDataMissing):
            bl.build_invoice(ctx(), "HEALTH_PLAN", None)
        no_bank = dict(LETTERS, remit={})
        with self.assertRaises(bl.LetterDataMissing):
            bl.build_invoice(ctx(letters=no_bank, invoices=[self.inv("PROVIDER", 100)]), "PROVIDER", None)


class StageLetters(unittest.TestCase):
    def test_every_stage_letter_builds(self):
        c = ctx(eligibility={"result": "INELIGIBLE", "reason": "below_threshold", "evidence": {"threshold_min_cents": 2500000}},
                rfi={"side": "PROVIDER", "items": ["Itemized bill", "Medical records"], "message": "", "due_date": "2026-09-16"},
                details={"withdrawal_dismissed_reason": "Provider No Response"})
        for key, fn in bl.BUILDERS.items():
            if key in ("final_order_rationale", "cover_letter_final_notice"):
                continue
            w = words(fn(c, None))
            self.assertIn("FL26-000216", w, key)
        self.assertIn("$25,000.00", words(bl.build_ineligible(c, None)))
        self.assertIn("Medical records", words(bl.build_rfi(c, None)))
        self.assertIn("DEFAULT DETERMINATION RECOMMENDATION", words(bl.build_plan_notification(c, None)))

    def test_ineligible_needs_an_ineligible_review(self):
        with self.assertRaises(bl.LetterDataMissing):
            bl.build_ineligible(ctx(eligibility={"result": "ELIGIBLE", "reason": "", "evidence": {}}), None)

    def test_counting(self):
        self.assertEqual(bl.counted(7, "claim"), "seven (7) claims")
        self.assertEqual(bl.counted(1, "patient"), "one (1) patient")
        self.assertEqual(bl.long_date("2025-03-03"), "March 3, 2025")


if __name__ == "__main__":
    unittest.main()
